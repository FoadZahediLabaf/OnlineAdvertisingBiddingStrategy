"""Build a short 4-minute presentation deck for Requirement 3."""

from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "slides" / "Requirement3_PrimalDual_4min_presentation.pptx"
ASSET = ROOT / "slides" / "generated_current"
SUMMARY = ROOT / "results" / "processed" / "req3_on_req4_shared" / "req3_factorized_summary.csv"

NAVY = RGBColor(13, 39, 63)
MUTED = RGBColor(95, 112, 137)
BLUE = RGBColor(48, 109, 246)
GREEN = RGBColor(18, 137, 110)
ORANGE = RGBColor(235, 111, 20)
PURPLE = RGBColor(124, 58, 237)
PALE = RGBColor(244, 247, 251)
LINE = RGBColor(214, 224, 236)
WHITE = RGBColor(255, 255, 255)
BLACK = RGBColor(26, 32, 44)


def set_text(shape, text, size=18, bold=False, color=BLACK, align=None):
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    if align is not None:
        p.alignment = align
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return shape


def add_text(slide, x, y, w, h, text, size=18, bold=False, color=BLACK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    return set_text(box, text, size=size, bold=bold, color=color, align=align)


def add_title(slide, title, subtitle=None):
    add_text(slide, 0.6, 0.38, 12.0, 0.55, title, size=30, bold=True, color=NAVY)
    if subtitle:
        add_text(slide, 0.62, 0.98, 11.8, 0.35, subtitle, size=13, color=MUTED)


def add_bar(slide):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    shape.fill.solid()
    shape.fill.fore_color.rgb = BLUE
    shape.line.fill.background()


def add_card(slide, x, y, w, h, title, value, note, color=BLUE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = WHITE
    card.line.color.rgb = LINE
    card.line.width = Pt(1.6)
    add_text(slide, x + 0.2, y + 0.18, w - 0.4, 0.25, title.upper(), size=10, bold=True, color=color)
    add_text(slide, x + 0.2, y + 0.55, w - 0.4, 0.45, value, size=23, bold=True, color=NAVY)
    add_text(slide, x + 0.2, y + 1.12, w - 0.4, 0.35, note, size=10, color=MUTED)


def add_bullet_list(slide, x, y, w, h, items, size=16):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.color.rgb = BLACK
        p.space_after = Pt(7)
    return box


def add_code_box(slide, x, y, w, h, lines):
    rect = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    rect.fill.solid()
    rect.fill.fore_color.rgb = RGBColor(248, 250, 252)
    rect.line.color.rgb = LINE
    rect.line.width = Pt(1.3)
    box = slide.shapes.add_textbox(Inches(x + 0.25), Inches(y + 0.18), Inches(w - 0.5), Inches(h - 0.3))
    tf = box.text_frame
    tf.clear()
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Consolas"
        p.font.size = Pt(12)
        p.font.color.rgb = BLACK
        p.space_after = Pt(2)
    return box


def add_picture(slide, name, x, y, w, h):
    slide.shapes.add_picture(str(ASSET / name), Inches(x), Inches(y), width=Inches(w), height=Inches(h))


def add_footer(slide, number):
    add_text(slide, 12.25, 7.02, 0.45, 0.2, str(number), size=8, color=MUTED, align=PP_ALIGN.RIGHT)


def add_mini_conflict_graph(slide, x, y):
    points = {
        "0": (x + 1.15, y + 0.10),
        "1": (x + 2.05, y + 0.70),
        "2": (x + 1.70, y + 1.70),
        "3": (x + 0.60, y + 1.70),
        "4": (x + 0.25, y + 0.70),
    }
    for a, b in [("0", "1"), ("1", "2"), ("2", "3"), ("3", "4"), ("4", "0")]:
        x1, y1 = points[a]
        x2, y2 = points[b]
        edge = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(x1),
            Inches(y1),
            Inches(x2),
            Inches(y2),
        )
        edge.line.color.rgb = ORANGE
        edge.line.width = Pt(4.5)
    for label, (px, py) in points.items():
        node = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(px - 0.18), Inches(py - 0.18), Inches(0.36), Inches(0.36))
        node.fill.solid()
        node.fill.fore_color.rgb = NAVY
        node.line.color.rgb = WHITE
        set_text(node, label, size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    add_text(slide, x - 0.05, y + 2.02, 2.6, 0.3, "Edges = incompatible pairs", size=10, color=MUTED, align=PP_ALIGN.CENTER)


def build():
    summary = pd.read_csv(SUMMARY)
    max_cost = summary["cost_mean"].max()
    min_cost = summary["cost_mean"].min()
    low = summary[summary["skl_label"] == "low_skl"]["utility_mean"].mean()
    high = summary[summary["skl_label"] == "high_skl"]["utility_mean"].mean()

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # Slide 1
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    add_bar(slide)
    add_title(
        slide,
        "Requirement 3: primal-dual bidding",
        "Multiple campaigns, global budget, full-information feedback.",
    )
    add_card(slide, 0.75, 1.65, 2.8, 1.55, "Horizon", "T=3000", "10 shared realizations per cell", BLUE)
    add_card(slide, 3.85, 1.65, 2.8, 1.55, "Budget", "B=2500", "rho = B / T per round", GREEN)
    add_card(slide, 6.95, 1.65, 2.8, 1.55, "Campaigns", "N=5", "5-cycle conflict graph", ORANGE)
    add_card(slide, 10.05, 1.65, 2.55, 1.55, "Grid", "3 x 3", "sKL x frequency", PURPLE)
    add_bullet_list(
        slide,
        0.9,
        4.05,
        8.1,
        1.4,
        [
            "Goal: learn which feasible super-arm and bids produce high utility while respecting the total budget.",
            "Feasible actions are 5 singletons plus 5 non-adjacent pairs from the conflict graph.",
            "The same precomputed market realizations are reused for fair comparison with the other algorithms.",
        ],
        size=16,
    )
    add_mini_conflict_graph(slide, 9.55, 4.05)
    add_footer(slide, 1)

    # Slide 2
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    add_bar(slide)
    add_title(slide, "Algorithm: factorized primal-dual Hedge", "High-level pseudocode included as requested in the project deliverable.")
    add_code_box(
        slide,
        0.65,
        1.45,
        6.1,
        5.25,
        [
            "Initialize q over feasible super-arms",
            "Initialize p_i over bid levels for each campaign i",
            "Initialize dual multiplier lambda = 0",
            "",
            "for t = 1,...,T:",
            "    sample a feasible super-arm S_t ~ q",
            "    sample one affordable bid for each i in S_t",
            "    set inactive campaign bids to 0",
            "    observe full market vector m_t",
            "    score every campaign-bid arm:",
            "        ell_t,i(g) = utility_t,i(g) - lambda * cost_t,i(g)",
            "    update Hedge weights with normalized scores",
            "    update lambda by projected gradient on spent budget",
        ],
    )
    add_bullet_list(
        slide,
        7.15,
        1.65,
        5.2,
        3.9,
        [
            "Primal side: Hedge increases probability of high-score super-arms and bid levels.",
            "Dual side: lambda penalizes expensive choices when cumulative spending is too high.",
            "Budget safety: a hard filter removes bids that cannot fit in the remaining budget.",
            "Full feedback lets the learner evaluate all campaign-bid arms after each round.",
        ],
        size=16,
    )
    add_footer(slide, 2)

    # Slide 3
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    add_bar(slide)
    add_title(slide, "Shared environment and benchmarks", "The experiment separates market generation from the learner.")
    add_picture(slide, "regime_schedule_overview.png", 0.65, 1.42, 5.95, 2.25)
    add_picture(slide, "opt_a_regret_comparison.png", 0.65, 4.0, 5.95, 2.35)
    add_bullet_list(
        slide,
        7.1,
        1.65,
        5.2,
        4.6,
        [
            "Market inputs are fixed first: phase ids, Gaussian means/covariances, change points, and sampled highest competing bids.",
            "Main regret baseline: a piecewise hindsight oracle using the best feasible action for each stationary segment.",
            "Additional OPT-A baseline: one LP solved after the run to find the best fixed randomized action under average budget.",
            "Low frequency has only 2 well-separated changes; medium has 12; high has 131.",
        ],
        size=15,
    )
    add_footer(slide, 3)

    # Slide 4
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    add_bar(slide)
    add_title(slide, "Empirical results on the 3 x 3 grid", "Mean over 10 shared realizations in each cell.")
    add_picture(slide, "sensitivity_grid.png", 0.55, 1.32, 7.75, 4.7)
    add_card(slide, 8.75, 1.55, 3.45, 1.25, "Budget feasibility", f"max cost {max_cost:,.1f}", "No run exceeds B=2500", GREEN)
    add_card(slide, 8.75, 3.05, 3.45, 1.25, "sKL effect", f"{low:,.0f} -> {high:,.0f}", "Mean utility drops as sKL rises", ORANGE)
    add_card(slide, 8.75, 4.55, 3.45, 1.25, "Main pattern", "separation dominates", "Frequency changes regret less visibly", PURPLE)
    add_footer(slide, 4)

    # Slide 5
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    add_bar(slide)
    add_title(slide, "Unexpected behavior and interpretation", "The plots show where the learner struggles.")
    add_picture(slide, "diagnostic_low_frequency_medium_skl.png", 0.55, 1.35, 6.4, 4.8)
    add_bullet_list(
        slide,
        7.25,
        1.55,
        5.2,
        4.8,
        [
            "Regret is not sublinear against the strong piecewise hindsight oracle in harder settings.",
            "The learner is budget-feasible but sometimes conservative, especially when high sKL makes regimes more separated.",
            "Dual multiplier moves only when spending pressure approaches the budget pace; otherwise the hard filter and low costs keep it near zero.",
            "This is useful diagnostically: the limitation is learning/adaptation, not budget violation.",
        ],
        size=15,
    )
    add_footer(slide, 5)

    # Slide 6
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PALE
    add_bar(slide)
    add_title(slide, "Takeaways", "What Requirement 3 contributes to the final comparison.")
    add_card(slide, 0.85, 1.65, 3.55, 1.55, "Algorithm", "primal-dual Hedge", "Super-arm learning + bid learning", BLUE)
    add_card(slide, 4.85, 1.65, 3.55, 1.55, "Constraint", "0 violations", "Budget enforced in all 90 runs", GREEN)
    add_card(slide, 8.85, 1.65, 3.55, 1.55, "Comparison", "shared inputs", "Same matrices for other learners", ORANGE)
    add_bullet_list(
        slide,
        1.0,
        4.0,
        11.2,
        1.65,
        [
            "The implementation satisfies the full-feedback primal-dual requirement and records reproducible shared realizations.",
            "The strongest empirical message is that sKL controls difficulty more than change frequency in this grid.",
            "For the final project comparison, the important common object is the same market matrix and the same oracle definition.",
        ],
        size=16,
    )
    add_footer(slide, 6)

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
